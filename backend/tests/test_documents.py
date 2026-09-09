from io import BytesIO

import pymupdf
import pytest
from docx import Document

from app.documents.errors import (
    DocumentEmpty,
    DocumentHasNoExtractableText,
    DocumentParseFailed,
    DocumentStorageUnavailable,
    DocumentTooLarge,
    UnsupportedDocumentType,
)
from app.documents.service import DocumentService
from app.documents.storage import LocalFileStore


@pytest.fixture
def documents():
    return DocumentService(1024 * 1024, 100000)


async def test_markdown_structure_and_text(documents):
    source = documents.validate(
        "../../lecture.md",
        "text/markdown",
        b"# Vectors\nA vector.\n\n- Direction\n## Length\nMagnitude.",
    )
    assert source.filename == "lecture.md"
    parsed = await documents.parse(source)
    assert [s.heading for s in parsed.sections] == ["Vectors", "Length"]
    assert [s.order for s in parsed.sections] == [0, 1]
    assert "- Direction" in parsed.sections[0].text
    plain = await documents.parse(documents.validate("a.txt", "text/plain", b"First\n\nSecond"))
    assert plain.sections[0].text == "First\n\nSecond"


async def test_pdf_pages_and_no_text(documents):
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((72, 72), "Vectors have magnitude and direction.")
        pdf.new_page().insert_text((72, 72), "A matrix transforms a vector.")
        content = pdf.tobytes()
    parsed = await documents.parse(documents.validate("a.pdf", "application/pdf", content))
    assert parsed.metadata.page_count == 2
    assert [s.source_page_start for s in parsed.sections] == [1, 2]
    with pymupdf.open() as pdf:
        pdf.new_page()
        blank = pdf.tobytes()
    with pytest.raises(DocumentHasNoExtractableText):
        await documents.parse(documents.validate("blank.pdf", "application/pdf", blank))


async def test_docx_headings(documents):
    doc = Document()
    doc.add_heading("Eigenvalues", 1)
    doc.add_paragraph("A scalar describes the scaling.")
    doc.add_heading("Example", 2)
    doc.add_paragraph("A diagonal matrix.", style="List Bullet")
    buffer = BytesIO()
    doc.save(buffer)
    parsed = await documents.parse(
        documents.validate("a.docx", "application/octet-stream", buffer.getvalue())
    )
    assert [s.heading for s in parsed.sections] == ["Eigenvalues", "Example"]
    assert parsed.sections[1].text.startswith("- ")


@pytest.mark.parametrize(
    "name,mime,body,error",
    [
        ("a.exe", "application/octet-stream", b"MZcode", UnsupportedDocumentType),
        ("a.txt", "application/pdf", b"text", UnsupportedDocumentType),
        ("a.pdf", "application/pdf", b"not pdf", UnsupportedDocumentType),
        ("a.txt", "text/plain", b"", DocumentEmpty),
        ("a.txt", "text/plain", b"MZ\x00", UnsupportedDocumentType),
        ("a.txt", "text/plain", b"a" * (1024 * 1024 + 1), DocumentTooLarge),
    ],
    ids=["executable", "mime-mismatch", "signature", "empty", "binary", "oversized"],
)
def test_validation(documents, name, mime, body, error):
    with pytest.raises(error):
        documents.validate(name, mime, body)


async def test_empty_and_corrupt(documents):
    with pytest.raises(DocumentEmpty):
        await documents.parse(documents.validate("a.txt", "text/plain", b" \n "))
    with pytest.raises(DocumentParseFailed):
        await documents.parse(documents.validate("a.pdf", "application/pdf", b"%PDF-broken"))


async def test_storage(tmp_path):
    store = LocalFileStore(tmp_path)
    key = await store.save(b"private")
    assert await store.read(key) == b"private"
    with pytest.raises(DocumentStorageUnavailable):
        await store.read("../outside")
    await store.delete(key)
    assert not (tmp_path / key).exists()
